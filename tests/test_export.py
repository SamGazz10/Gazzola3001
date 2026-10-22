"""Checks on the spreadsheet export (small run)."""
from openpyxl import load_workbook

from overtake.export import ATTEMPT_COLUMNS, write_workbook


def test_workbook_has_every_sheet_and_attempt(tmp_path):
    """A small workbook has the expected sheets, one row per attempt, all three splits and valid probabilities."""
    path = write_workbook(tmp_path / "small.xlsx", n=400, log=lambda _: None)
    wb = load_workbook(path, read_only=True)
    assert wb.sheetnames[:6] == ["About", "Summary", "Metrics", "Inputs", "Model parameters", "Attempts"]
    rows = list(wb["Attempts"].iter_rows(values_only=True))
    assert list(rows[0]) == [name for name, _, _ in ATTEMPT_COLUMNS]
    assert len(rows) == 401
    split = rows[0].index("split")
    assert {r[split] for r in rows[1:]} == {"train", "validation", "test"}
    for column in ("p_true", "p_network", "p_gradient_boosting", "p_logistic_regression"):
        i = rows[0].index(column)
        assert all(0 <= r[i] <= 1 for r in rows[1:])
