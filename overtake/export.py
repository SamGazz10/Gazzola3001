"""Export every simulated attempt and every experiment result to one Excel workbook.

    python -m overtake.export                         # writes data/overtake_simulations.xlsx
    python -m overtake.export --out my_copy.xlsx

The workbook is regenerated from the fixed seed rather than stored in git
(about 35 MB); the same code and seed always give the same numbers. It needs
the trained model in ``models/`` and the experiment results in
``reports/results/`` (``python -m overtake.train`` and
``python -m overtake.experiments all``).

| Sheet | Contents |
|---|---|
| About | What the workbook is, how it was made, a dictionary of the Attempts columns |
| Summary | Key figures recomputed with formulas from the Attempts sheet, beside the recorded values |
| Metrics | Every metric used to evaluate the project: definition, how it is computed, results |
| Inputs | The 19 inputs (factors) the model takes into account |
| Model parameters | Every constant in the physics and outcome models, with its source |
| Attempts | All 300,000 simulated attempts: inputs, physics results, outcome, true and predicted probability |
| Speed trace | Time history of both cars along the straight for two scenarios |
| Solver, Time step, Quadrature, Network selection, Ablation, Model comparison, Sensitivity, Performance | One sheet per experiment |
"""
import argparse
import datetime
import json
import subprocess
import time
from dataclasses import fields
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList

from . import baselines, physics, simulator
from .features import BINARY, CATEGORICAL, DEFAULTS, LABEL, NUMERIC, NUMERIC_RANGES, VOCABS
from .predict import OvertakePredictor
from .train import REPORTS, ROOT, split

OUT = ROOT / "data" / "overtake_simulations.xlsx"
SEED = 0

FONT = "Arial"
HEADER_FONT = Font(name=FONT, sz=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1C5CAB")
TITLE_FONT = Font(name=FONT, sz=14, bold=True)
SECTION_FONT = Font(name=FONT, sz=11, bold=True)
NOTE_FONT = Font(name=FONT, sz=10, italic=True, color="52514E")
WRAP = Alignment(wrap_text=True, vertical="top")

# What each input means and how the simulator uses it.
INPUTS = {
    "track_condition": ("Track", "category", "-", "Sets tyre grip for every compound, average speeds, how much "
                        "extra power becomes speed, execution noise, caution and completion odds",
                        "Surface: dry, damp or wet"),
    "segment": ("Track", "category", "-", "Completion odds (corner type at the end of the run)",
                "Where the move is completed: at the end of a straight, or into a slow, medium or fast corner"),
    "circuit": ("Track", "category", "-", "Completion odds (passability of the circuit)",
                "Circuit; 'generic' when unknown"),
    "attacker_tyre": ("Attacker", "category", "-", "Tyre grip: traction, corner-exit speed, braking duel",
                      "Attacker's tyre compound"),
    "defender_tyre": ("Defender", "category", "-", "Tyre grip: traction, corner-exit speed, braking duel",
                      "Defender's tyre compound"),
    "gap_s": ("Both", "number", "s", "Starting separation; slipstream strength; dirty-air grip loss",
              "Gap to the car ahead at the exit of the corner before the run"),
    "run_length_m": ("Track", "number", "m", "Distance integrated; how long power and energy differences act",
                     "Distance from that corner exit to the braking point"),
    "attacker_battery": ("Attacker", "number", "%", "Electrical energy available (100 % = 4 MJ)",
                         "Attacker's battery state of charge"),
    "defender_battery": ("Defender", "number", "%", "Electrical energy available; ability to fight back",
                         "Defender's battery state of charge"),
    "attacker_tyre_age": ("Attacker", "number", "laps", "Tyre wear and grip loss (with a cliff)",
                          "Laps on the attacker's tyres"),
    "defender_tyre_age": ("Defender", "number", "laps", "Tyre wear and grip loss (with a cliff)",
                          "Laps on the defender's tyres"),
    "track_temp_c": ("Track", "number", "deg C", "Tyre wear rate; hard tyres lose grip when cold",
                     "Track temperature"),
    "attacker_skill": ("Attacker", "number", "0-100", "Execution consistency; braking duel",
                       "Attacker's racecraft rating"),
    "defender_skill": ("Defender", "number", "0-100", "Defensive positioning; braking duel",
                       "Defender's defensive rating"),
    "top_speed_delta_kph": ("Attacker", "number", "km/h", "Attacker's drag area (lower drag = higher top speed)",
                            "Attacker car's straight-line speed advantage, energy aside"),
    "pace_delta_s": ("Attacker", "number", "s/lap", "Attacker's corner-exit speed",
                     "Attacker car's lap-time advantage, tyres aside"),
    "attacker_boost": ("Attacker", "true/false", "-", "Full MGU-K deployment (otherwise 55 % on the standard map)",
                       "Attacker uses Boost for the attack"),
    "defender_boost": ("Defender", "true/false", "-", "Full MGU-K deployment; fight-back on the exit",
                       "Defender uses Boost to defend"),
    "overtake_mode": ("Attacker", "true/false", "-", "+0.5 MJ and full power up to 337 km/h instead of 290 km/h",
                      "Attacker was within 1 s at the detection point"),
}

ATTEMPT_COLUMNS = [
    ("attempt_id", "-", "Row number of the simulated attempt"),
    ("split", "-", "train (80 %), validation (10 %) or test (10 %)"),
    *[(c, "1 = yes, 0 = no" if c in BINARY else INPUTS[c][2], INPUTS[c][4]) for c in INPUTS],
    ("attacker_grip", "-", "Attacker tyre grip relative to a new soft on a dry track"),
    ("defender_grip", "-", "Defender tyre grip relative to a new soft on a dry track"),
    ("deficit_at_braking_point_s", "s", "Physics result: gap still to close when the defender reaches the "
     "braking point (negative = attacker already ahead)"),
    ("attacker_speed_kph", "km/h", "Physics result: attacker speed at the braking point"),
    ("defender_speed_kph", "km/h", "Physics result: defender speed at the braking point"),
    ("attacker_energy_used_mj", "MJ", "Physics result: attacker battery energy used along the run"),
    ("defender_energy_used_mj", "MJ", "Physics result: defender battery energy used along the run"),
    ("run_time_s", "s", "Physics result: time for the defender to cover the run"),
    ("effective_deficit_s", "s", "Deficit plus wet-weather caution and the defender's positioning"),
    ("execution_sd_s", "s", "Standard deviation of the attacker's execution noise"),
    ("completion_logit_z0", "-", "Log-odds that a lunge sticks with zero margin"),
    ("overtake", "0/1", "Simulated outcome: 1 if the overtake stuck"),
    ("p_true", "probability", "Exact probability of an overtake (physics + Gauss-Legendre quadrature)"),
    ("p_network", "probability", "Prediction of the final neural network (ensemble of five)"),
    ("p_gradient_boosting", "probability", "Prediction of gradient-boosted trees (trained on the train split)"),
    ("p_logistic_regression", "probability", "Prediction of logistic regression (trained on the train split)"),
]


def _git_commit():
    """Short hash of the checked-out commit, recorded in the workbook so a copy can be traced to its code.

    Returns ``"unknown"`` when git is missing or the folder is not a repository.
    """
    try:
        return subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT, capture_output=True,
                              text=True, check=True).stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        return "unknown"


def _results(name):
    """Load ``reports/results/<name>.json`` (raises FileNotFoundError if the experiment has not run)."""
    return json.loads((REPORTS / "results" / f"{name}.json").read_text())


def attempts_table(n=300_000, seed=SEED, log=print):
    """Every simulated attempt with its physics results, outcome and predictions.

    Parameters
    ----------
    n : int, optional
        Number of attempts (300,000 is the dataset used throughout the project).
    seed : int, optional
    log : callable, optional

    Returns
    -------
    pandas.DataFrame
        One row per attempt, columns as listed in `ATTEMPT_COLUMNS`.
    """
    log(f"simulating {n:,} attempts")
    df = simulator.generate_dataset(n, seed=seed)
    train_df, val_df, test_df = split(df, seed)
    labels = pd.Series("train", index=df.index)
    labels[val_df.index] = "validation"
    labels[test_df.index] = "test"

    log("recording physics results")
    parts = []
    for start in range(0, n, 20_000):
        part = df.iloc[start:start + 20_000]
        run, grip_a, grip_d = simulator.closing_run(part)
        deficit, sigma, z0 = simulator.outcome_parameters(part, run, grip_a, grip_d)
        parts.append(pd.DataFrame({
            "attacker_grip": grip_a, "defender_grip": grip_d,
            "deficit_at_braking_point_s": run.deficit_s,
            "attacker_speed_kph": run.attacker_speed * 3.6, "defender_speed_kph": run.defender_speed * 3.6,
            "attacker_energy_used_mj": run.attacker_energy_j / 1e6,
            "defender_energy_used_mj": run.defender_energy_j / 1e6,
            "run_time_s": run.time_s, "effective_deficit_s": deficit, "execution_sd_s": sigma,
            "completion_logit_z0": z0,
        }, index=part.index))
    details = pd.concat(parts)

    log("predicting with the network and the alternatives")
    predictions = pd.DataFrame({
        "p_network": OvertakePredictor.load().predict(df).astype(np.float64),  # float32 would round untidily
        "p_gradient_boosting": baselines.gradient_boosting(seed).fit(train_df).predict(df),
        "p_logistic_regression": baselines.logistic_regression().fit(train_df).predict(df),
    }, index=df.index)

    out = pd.concat([df[list(INPUTS)], details, df[[LABEL, "p_true"]], predictions], axis=1)
    out.insert(0, "split", labels)
    out.insert(0, "attempt_id", np.arange(1, n + 1))
    decimals = {"gap_s": 4, "run_length_m": 1, "attacker_battery": 2, "defender_battery": 2, "track_temp_c": 1,
                "attacker_skill": 1, "defender_skill": 1, "top_speed_delta_kph": 2, "pace_delta_s": 3,
                "attacker_grip": 4, "defender_grip": 4, "deficit_at_braking_point_s": 4, "attacker_speed_kph": 1,
                "defender_speed_kph": 1, "attacker_energy_used_mj": 3, "defender_energy_used_mj": 3,
                "run_time_s": 3, "effective_deficit_s": 4, "execution_sd_s": 3, "completion_logit_z0": 3,
                "p_true": 5, "p_network": 5, "p_gradient_boosting": 5, "p_logistic_regression": 5}
    out = out.round(decimals)
    # Flags as 1/0 like the outcome (LibreOffice would turn TRUE/FALSE cells into formulas).
    for col in ("attacker_tyre_age", "defender_tyre_age", LABEL, *BINARY):
        out[col] = out[col].astype(int)
    return out[[c for c, _, _ in ATTEMPT_COLUMNS]]


class _Sheet:
    """Small helper for appending styled rows to a write-only worksheet."""

    def __init__(self, wb, name, widths):
        """Create sheet ``name`` in workbook ``wb`` with the given column widths (characters)."""
        self.ws = wb.create_sheet(name)
        self.rows = 0  # rows written so far; write-only sheets do not track this
        for i, width in enumerate(widths, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = width

    def append(self, cells):
        """Append one row of cells and count it."""
        self.ws.append(cells)
        self.rows += 1

    def _cell(self, value, font=None, fill=None, wrap=False, number_format=None):
        """A write-only cell holding ``value`` with the optional font, fill, wrapping and number format."""
        cell = WriteOnlyCell(self.ws, value=value)
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if wrap:
            cell.alignment = WRAP
        if number_format:
            cell.number_format = number_format
        return cell

    def title(self, text):
        """Write the sheet's title line."""
        self.append([self._cell(text, TITLE_FONT)])

    def section(self, text):
        """Write a blank line, then a section heading."""
        self.append([])
        self.append([self._cell(text, SECTION_FONT)])

    def note(self, text):
        """Write an explanatory note in the muted note style."""
        self.append([self._cell(text, NOTE_FONT)])

    def header(self, names):
        """Write a table header row (bold, shaded, wrapped)."""
        self.append([self._cell(n, HEADER_FONT, HEADER_FILL, wrap=True) for n in names])

    def row(self, values, wrap=False, formats=None):
        """Write one data row; ``formats`` maps a column index to an Excel number format."""
        formats = formats or {}
        self.append([self._cell(_native(v), wrap=wrap, number_format=formats.get(i))
                     for i, v in enumerate(values)])

    def table(self, header, rows, wrap=False, formats=None):
        """Write a header row followed by every row in ``rows``."""
        self.header(header)
        for r in rows:
            self.row(r, wrap=wrap, formats=formats)


def _native(value):
    """Convert a NumPy scalar to the plain Python type openpyxl accepts, and a list to comma-separated text."""
    if isinstance(value, (np.bool_,)):
        return bool(value)
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    return value


def _metric_rows():
    """Every metric used in the project, with definition, computation and results."""
    c, sel, ab = _results("comparison"), _results("selection"), _results("ablation")
    ts, sol, quad = _results("timestep"), _results("solver"), _results("quadrature")
    perf, sens = _results("performance"), _results("sensitivity")
    nn, gb, lr = (c["models"][m] for m in ("Neural network", "Gradient boosting", "Logistic regression"))
    q = c["error_quantiles"]
    pts = lambda x: round(100 * x, 2)  # noqa: E731
    i005 = ts["dt"].index(0.05)
    raw = ab["raw inputs only"]
    fl = perf["flops"]

    def per_model(key, scale=1.0, nd=4):
        """The metric ``key`` for the network, gradient boosting and logistic regression, scaled and rounded."""
        return [round(scale * m[key], nd) for m in (nn, gb, lr)]

    rows = [
        # category, metric, definition, how computed, units, better, where, NN, GB, LR, other
        ("Prediction accuracy", "Mean absolute error vs true probability",
         "Average |predicted P - true P| over the test set", "mean(|p - p_true|)", "percentage points",
         "lower", "train.metrics", *per_model("mae_vs_true", 100, 2), ""),
        ("Prediction accuracy", "Maximum absolute error vs true probability", "Worst single test scenario",
         "max(|p - p_true|)", "percentage points", "lower", "train.metrics", *per_model("max_error_vs_true", 100, 1), ""),
        ("Prediction accuracy", "Median absolute error vs truth", "50th percentile of |p - p_true|", "quantile 0.5",
         "percentage points", "lower", "experiments.comparison",
         *[pts(q[m]["q50"]) for m in ("Neural network", "Gradient boosting", "Logistic regression")], ""),
        ("Prediction accuracy", "90th / 99th / 99.9th percentile error", "Tail of |p - p_true|", "quantiles",
         "percentage points", "lower", "experiments.comparison",
         *[f"{pts(q[m]['q90'])} / {pts(q[m]['q99'])} / {pts(q[m]['q99.9'])}"
           for m in ("Neural network", "Gradient boosting", "Logistic regression")], ""),
        ("Prediction accuracy", "Log loss (binary cross-entropy)",
         "How well probabilities explain the 0/1 outcomes; also the training loss",
         "-mean(y log p + (1-y) log(1-p))", "nats", "lower", "train.metrics", *per_model("log_loss"), ""),
        ("Prediction accuracy", "Brier score", "Mean squared error of the probability against the 0/1 outcome",
         "mean((p - y)^2)", "-", "lower", "train.metrics", *per_model("brier"), ""),
        ("Prediction accuracy", "ROC AUC", "Probability a random overtake is ranked above a random failure",
         "area under the ROC curve", "-", "higher", "train.metrics", *per_model("roc_auc"), ""),
        ("Prediction accuracy", "Accuracy at 0.5", "Share of outcomes on the right side of 50 %",
         "mean((p >= 0.5) == y)", "-", "higher", "train.metrics", *per_model("accuracy"), ""),
        ("Calibration", "Expected calibration error",
         "Average gap between predicted probability and observed rate, 10 equal-width bins",
         "sum over bins of share x |mean p - mean y|", "-", "lower", "train.expected_calibration_error",
         *per_model("calibration_error"), ""),
        ("Robustness", "Error by operating region", "Mean error vs truth per track condition, segment and "
         "tyre suitability", "mean(|p - p_true|) within each region", "percentage points", "lower",
         "experiments.comparison", "see Model comparison sheet", "", "", ""),
        ("Robustness", "Tyre-mismatch share of the worst errors",
         "Share of the worst 1 % / 0.1 % of errors where a car's tyres do not suit the surface "
         f"({round(100 * c['tyre_mismatch_share'])} % of all test scenarios)", "share", "%", "-",
         "experiments.comparison",
         f"{round(100 * q['Neural network']['tyre_mismatch_share_of_worst_1pct'])} / "
         f"{round(100 * q['Neural network']['tyre_mismatch_share_of_worst_0.1pct'])}", "", "", ""),
        ("Robustness", "Extrapolation error", "Mean error vs truth on straights of 1300-2000 m (beyond training)",
         "mean(|p - p_true|)", "percentage points", "lower", "experiments.comparison",
         *[pts(c["extrapolation_1300_2000m"][m]) for m in ("Neural network", "Gradient boosting",
                                                          "Logistic regression")], ""),
        ("Robustness", "Error at the default scenario", "Prediction minus truth for the default scenario",
         "p - p_true", "percentage points", "closer to 0", "experiments.sensitivity",
         pts(sens["default_scenario"]["network"] - sens["default_scenario"]["exact"]), "", "", ""),
        ("Robustness", "Worst error along the battery axis",
         "Largest |p - p_true| as the attacker battery goes 30-100 % at the default scenario (validation, "
         "ensemble of 5)", "max |p - p_true|", "percentage points", "lower", "experiments.ablation",
         pts(raw["ensemble_of_5"]["battery_axis_max_error"]), "", "", ""),
        ("Robustness", "Run-to-run (seed) variation",
         "Standard deviation across five training seeds of a single network's default-scenario error",
         "sd over seeds", "percentage points", "lower", "experiments.ablation",
         pts(raw["single_model"]["default_scenario_error"]["sd"]), "", "", ""),
        ("Robustness", "Sensitivity agreement", "Network vs simulator probability at each of 30 one-at-a-time "
         "settings around the default", "|p - p_true| per setting", "percentage points", "lower",
         "experiments.sensitivity", "see Sensitivity sheet", "", "", ""),
        ("Model selection", "Validation log loss ± standard error",
         "Log loss on the validation set and its standard error; the one-standard-error rule picks the "
         "smallest network within one standard error of the lowest", "mean and sd/sqrt(n) of per-attempt loss",
         "nats", "lower", "experiments.selection", f"selected: {sel['selected']}", "", "",
         f"threshold {sel['threshold']:.4f}"),
        ("Model selection", "Validation error vs training-set size", "Mean error vs truth with 10k-240k "
         "training attempts", "mean(|p - p_true|)", "percentage points", "lower", "experiments.selection",
         " / ".join(str(pts(d["val_mae_vs_true"])) for d in sel["data_sizes"]), "", "",
         "sizes " + " / ".join(f"{d['train_size']:,}" for d in sel["data_sizes"])),
        ("Cost", "Training time", "Wall-clock time to fit (network: five networks)", "timer", "s", "lower",
         "experiments.comparison", *per_model("fit_seconds", 1, 1), "direct simulation: none"),
        ("Cost", "Prediction time", "Per scenario in a batch of 10,000, including input checks", "timer",
         "microseconds", "lower", "experiments.comparison", *per_model("predict_us_per_scenario", 1, 2),
         f"direct simulation: {c['models']['Direct simulation']['predict_us_per_scenario']:.0f}"),
        ("Cost", "Model size", "Bytes of parameters (network) or pickled model", "count", "kB", "lower",
         "experiments.comparison", *[round(m["model_bytes"] / 1e3, 1) for m in (nn, gb, lr)], ""),
        ("Cost", "FLOPs per prediction", "Floating-point operations (multiply-add = 2)",
         "from layer sizes (network); hand count x RK4 steps (simulation)", "FLOP", "lower",
         "experiments.performance", fl["network_flops_per_prediction"], "", "",
         f"direct simulation: ~{fl['simulation_flops_per_scenario_estimate']:,.0f}"),
        ("Cost", "Achieved FLOP rate", "FLOPs per prediction x predictions / time", "-", "GFLOP/s", "higher",
         "experiments.performance", round(fl["network_gflops_achieved"], 1), "", "",
         f"direct simulation: {fl['simulation_gflops_achieved']:.2f}"),
        ("Cost", "Arithmetic intensity", "FLOPs per byte moved for network inference, batch of 1 / 64 / 1024",
         "FLOPs / (weight bytes + input/output bytes)", "FLOP/byte", "-", "experiments.performance",
         " / ".join(f"{fl['network_arithmetic_intensity_flop_per_byte'][b]:.1f}" for b in ("1", "64", "1024")),
         "", "", ""),
        ("Cost", "Memory", "Network parameters; 300k-attempt dataset; ODE state per 20k chunk", "bytes", "MB",
         "lower", "experiments.performance", round(perf["memory"]["network_parameter_bytes"] / 1e6, 3), "", "",
         f"dataset {perf['memory']['dataset_300k_bytes'] / 1e6:.0f}; ODE state "
         f"{perf['memory']['ode_state_bytes_per_20k_chunk'] / 1e6:.2f}"),
        ("Cost", "Vectorisation speed-up", "200 scenarios integrated together vs one at a time", "time ratio",
         "x", "higher", "experiments.performance", "", "", "",
         round(perf["vectorisation_200_scenarios"]["speedup"], 1)),
        ("Numerical accuracy", "Error vs analytical solution",
         "Position error at 12.8 s for constant power / drag only", "|x - x_exact|", "m", "lower",
         "experiments.solver", "", "", "", "see Solver sheet"),
        ("Numerical accuracy", "Observed order of convergence",
         "Slope of log(error) vs log(time step); RK4 should be 4, Euler 1", "least-squares slope", "-",
         "matches theory", "experiments.solver", "", "", "",
         "RK4 " + " / ".join(f"{r['rk4_order']:.2f}" for r in sol.values()) + "; Euler "
         + " / ".join(f"{r['euler_order']:.2f}" for r in sol.values())),
        ("Numerical accuracy", "Braking-point deficit error vs fine-step reference",
         "Max and mean |deficit(dt) - deficit(0.001 s)| over 2,000 scenarios at dt = 0.05 s",
         "max/mean |difference|", "s", "lower (limit 0.001)", "experiments.timestep", "", "", "",
         f"RK4 max {ts['rk4']['max_error_s'][i005]:.1e}, mean {ts['rk4']['mean_error_s'][i005]:.1e}; "
         f"Euler max {ts['euler']['max_error_s'][i005]:.1e}"),
        ("Numerical accuracy", "Quadrature error", "Max |P(rule) - P(reference)| over 2,000 scenarios",
         "vs 200-node Gauss-Legendre", "probability", "lower", "experiments.quadrature", "", "", "",
         f"Gauss-Legendre 16 nodes: {quad['rules']['Gauss-Legendre']['max_error'][4]:.1e}"),
        ("Numerical accuracy", "Monte Carlo RMS error", "RMS of MC estimate - exact P, against theory "
         "sqrt(p(1-p)/N)", "sqrt(mean((P_mc - P)^2))", "probability", "lower", "experiments.quadrature", "",
         "", "", f"N=1600: {quad['monte_carlo']['rms_error'][2]:.4f} (theory {quad['monte_carlo']['theory'][2]:.4f})"),
        ("Physical plausibility", "Constraint checks", "Energy used <= energy available; speeds within "
         "physical bounds; P moves the expected way for 10 inputs", "automated tests", "-", "always pass",
         "tests/", "", "", "", "pass"),
    ]
    return rows


def _write_reference_sheets(wb, attempts):
    """Write the explanatory sheets: About, Summary, Metrics, Inputs and Model parameters.

    Summary formulas refer to ranges of the Attempts sheet, so they stay live in Excel.

    Parameters
    ----------
    wb : openpyxl.Workbook
        Write-only workbook.
    attempts : pandas.DataFrame
        Output of `attempts_table`; its column order fixes the Attempts sheet's columns.
    """
    c = _results("comparison")
    n = len(attempts)
    last = n + 1
    col = {name: get_column_letter(i) for i, name in enumerate(attempts.columns, start=1)}

    def rng(name):
        """Absolute Excel range of column ``name`` on the Attempts sheet, every data row."""
        return f"Attempts!${col[name]}$2:${col[name]}${last}"

    # About
    s = _Sheet(wb, "About", [34, 16, 100])
    s.title("F1 2026 overtake probability: simulation data")
    s.note("MMA3001 project. Generated by `python -m overtake.export`; do not edit by hand.")
    s.row(["Generated", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")])
    s.row(["Code version (git commit)", _git_commit()])
    s.row(["Random seed", SEED])
    s.row(["Simulated attempts", n])
    s.row(["Source of experiment results", "reports/results/*.json in the repository"])
    s.section("Sheets")
    s.table(["Sheet", "", "Contents"], [
        ("Summary", "", "Key figures recomputed with formulas from the Attempts sheet, next to the recorded values"),
        ("Metrics", "", "Every metric used to evaluate the project: definition, how it is computed, results"),
        ("Inputs", "", "The 19 inputs the model takes into account, with units, allowed values and their role"),
        ("Model parameters", "", "Every constant in the physics and outcome models, with its source"),
        ("Attempts", "", "All simulated attempts: inputs, physics results, outcome, true and predicted probability"),
        ("Speed trace", "", "Speed and battery energy of both cars along the straight, two scenarios"),
        ("Solver", "", "Euler and RK4 against analytical solutions"),
        ("Time step", "", "Full-model error and cost against time step, before and after the slipstream fix"),
        ("Quadrature", "", "Quadrature rules and Monte Carlo against a reference"),
        ("Network selection", "", "Network sizes and training-set sizes on the validation set"),
        ("Ablation", "", "Physics-informed inputs and ensembling over five seeds"),
        ("Model comparison", "", "Network vs gradient boosting vs logistic regression vs direct simulation"),
        ("Sensitivity", "", "One input at a time around the default scenario; wet-track check"),
        ("Performance", "", "Profiling, vectorisation, chunk size, FLOPs, memory"),
    ], wrap=True)
    s.section("Attempts columns")
    s.table(["Column", "Units", "Meaning"], ATTEMPT_COLUMNS, wrap=True)

    # Summary: formulas over the Attempts sheet, next to the values recorded by the experiments.
    s = _Sheet(wb, "Summary", [44, 18, 18, 18, 18, 50])
    s.title("Summary")
    s.note("Black figures are formulas over the Attempts sheet; 'recorded' figures come from "
           "reports/results/comparison.json, written by `python -m overtake.experiments comparison`.")
    s.section("Dataset")
    s.header(["Figure", "Value (formula)", "", "", "", "Note"])
    split_rng = rng("split")
    s.row(["Simulated attempts", f"=COUNTA({rng('attempt_id')})", "", "", "", ""])
    for name in ("train", "validation", "test"):
        s.row([f"{name.capitalize()} attempts", f'=COUNTIF({split_rng},"{name}")', "", "", "",
               "80 / 10 / 10 split" if name == "train" else ""])
    s.row(["Overtake rate (all attempts)", f"=AVERAGE({rng(LABEL)})", "", "", "", "share of attempts that stuck"],
          formats={1: "0.0%"})
    s.row(["Mean true probability (all attempts)", f"=AVERAGE({rng('p_true')})", "", "", "",
           "should match the overtake rate"], formats={1: "0.0%"})

    s.section("Test set, recomputed from the Attempts sheet")
    s.header(["Model", "Mean error vs truth (points)", "Brier score", "Accuracy at 0.5",
              "Recorded mean error (points)", "Difference (points)"])
    test_n = f'COUNTIF({split_rng},"test")'
    recorded = {"p_network": "Neural network", "p_gradient_boosting": "Gradient boosting",
                "p_logistic_regression": "Logistic regression"}
    for pcol, name in recorded.items():
        r = s.rows + 1  # the row this model's figures go on
        test = f'({split_rng}="test")'
        s.row([
            name,
            f"=100*SUMPRODUCT({test}*ABS({rng(pcol)}-{rng('p_true')}))/{test_n}",
            f"=SUMPRODUCT({test}*({rng(pcol)}-{rng(LABEL)})^2)/{test_n}",
            f"=SUMPRODUCT({test}*((({rng(pcol)}>=0.5)*1)={rng(LABEL)}))/{test_n}",
            round(100 * c["models"][name]["mae_vs_true"], 4),
            f"=B{r}-E{r}",
        ], formats={1: "0.00", 2: "0.0000", 3: "0.0%", 4: "0.00", 5: "0.0000"})
    s.note("Differences of about 0.0001 points come from rounding the stored probabilities to 5 decimals.")

    sens = _results("sensitivity")
    s.section("Default scenario (dry, new softs, full battery, 1 s behind on an 800 m straight)")
    s.header(["Figure", "Value", "", "", "", "Source"])
    s.row(["True probability (simulator)", round(sens["default_scenario"]["exact"], 4), "", "", "",
           "sensitivity.json"], formats={1: "0.0%"})
    s.row(["Neural network", round(sens["default_scenario"]["network"], 4), "", "", "", "sensitivity.json"],
          formats={1: "0.0%"})

    # Metrics
    s = _Sheet(wb, "Metrics", [20, 34, 50, 34, 16, 13, 26, 16, 16, 16, 34])
    s.title("Every metric used in the project")
    s.note("Results are on the 30,000-attempt test set unless the definition says otherwise. "
           "Recorded by the experiments in reports/results/*.json.")
    s.table(["Category", "Metric", "Definition", "How it is computed", "Units", "Better", "Code",
             "Neural network", "Gradient boosting", "Logistic regression", "Other / result"],
            _metric_rows(), wrap=True)

    # Inputs
    data_ranges = {}
    for name in NUMERIC:
        data_ranges[name] = f"{attempts[name].min():g} to {attempts[name].max():g}"
    s = _Sheet(wb, "Inputs", [22, 11, 11, 9, 34, 22, 12, 52, 48])
    s.title("The 19 inputs the model takes into account")
    s.note("'Allowed' is what the tool accepts (anything else raises an error); 'training range' is what the "
           "simulated data covers (outside it the tool warns that it is extrapolating).")
    rows = []
    for name, (car, kind, units, role, meaning) in INPUTS.items():
        if name in CATEGORICAL:
            allowed = ", ".join(VOCABS[CATEGORICAL[name]])
            seen = "all"
        elif name in NUMERIC:
            lo, hi = NUMERIC_RANGES[name]
            allowed = f"{lo:g} to {hi:g}"
            seen = data_ranges[name]
        else:
            allowed, seen = "true / false", "both"
        rows.append((name, car, kind, units, allowed, seen, _native(DEFAULTS[name]), role, meaning))
    s.table(["Input", "Car", "Type", "Units", "Allowed", "Training range", "Default", "How the simulator uses it",
             "Meaning"], rows, wrap=True)

    # Model parameters
    car = physics.CarParameters()
    units = {"mass_kg": "kg", "cda_m2": "m^2", "cla_m2": "m^2", "air_density": "kg/m^3", "rolling_coeff": "-",
             "ice_power_w": "W", "mguk_power_w": "W", "drivetrain_efficiency": "-", "mu_reference": "-",
             "driven_axle_share": "-", "tow_max": "-", "tow_length_m": "m", "car_length_m": "m",
             "energy_smoothing_j": "J"}
    sources = {"ice_power_w": "2026 rules as reported (~400 kW)", "mguk_power_w": "2026 rules (350 kW)",
               "mass_kg": "estimate: 2026 minimum ~768 kg plus fuel", "air_density": "standard air"}
    s = _Sheet(wb, "Model parameters", [34, 22, 10, 30, 60])
    s.title("Constants in the physics and outcome models")
    s.note("'Estimate' means an engineering judgement chosen for plausible results, not a value fitted to F1 data.")
    rows = [(f"CarParameters.{f.name}", getattr(car, f.name), units[f.name], "physics (both cars)",
             sources.get(f.name, "estimate")) for f in fields(car)]
    rows += [
        ("TAPER_STANDARD_KPH", "290 to 355", "km/h", "MGU-K power limit", "2026 rules as reported (Formula1.com)"),
        ("TAPER_OVERTAKE_MODE_KPH", "337 to 355", "km/h", "MGU-K limit with Overtake Mode",
         "2026 rules as reported (Formula1.com, Crash.net)"),
        ("BATTERY_USABLE_J", physics.BATTERY_USABLE_J, "J", "energy at 100 % charge",
         "estimate (storage capacity reported unchanged)"),
        ("OVERTAKE_MODE_EXTRA_J", physics.OVERTAKE_MODE_EXTRA_J, "J", "extra energy with Overtake Mode",
         "2026 rules as reported"),
        ("STANDARD_MAP_SHARE", simulator.STANDARD_MAP_SHARE, "-", "MGU-K share without Boost", "estimate"),
        ("EXIT_SPEED_MPS", simulator.EXIT_SPEED_MPS, "m/s", "corner-exit speed, new soft, dry", "estimate"),
        ("DIRTY_AIR_GRIP_LOSS", simulator.DIRTY_AIR_GRIP_LOSS, "-", "grip lost right behind a car", "estimate"),
        ("TOP_SPEED_REF_MPS", simulator.TOP_SPEED_REF_MPS, "m/s", "converts top-speed advantage to drag", "estimate"),
        ("LAP_TIME_REF_S", simulator.LAP_TIME_REF_S, "s", "converts pace advantage to exit speed", "estimate"),
        ("TIME_STEP_S", simulator.TIME_STEP_S, "s", "RK4 step", "time-step study"),
        ("LUNGE_RANGE_S", simulator.LUNGE_RANGE_S, "s", "deficit from which a lunge is possible", "estimate"),
        ("MAX_MARGIN_S", simulator.MAX_MARGIN_S, "s", "cap on the completion margin", "estimate"),
        ("MARGIN_LOGIT_PER_S", simulator.MARGIN_LOGIT_PER_S, "1/s", "completion log-odds per second of margin",
         "estimate"),
        ("FULL_BATTERY_SHARE", simulator.FULL_BATTERY_SHARE, "-", "share of sampled attempts at 100 % battery",
         "data design"),
        ("CLIFF_DEG_PER_LAP", simulator.CLIFF_DEG_PER_LAP, "grip/lap", "extra wear past the cliff", "estimate"),
    ]
    for table, label, unit in ((simulator.EXECUTION_SD, "EXECUTION_SD", "s"), (simulator.CAUTION_S, "CAUTION_S", "s"),
                               (simulator.CONDITION_LOGIT, "CONDITION_LOGIT", "log-odds"),
                               (simulator.SEGMENT_LOGIT, "SEGMENT_LOGIT", "log-odds"),
                               (simulator.DEG_PER_LAP, "DEG_PER_LAP", "grip/lap"),
                               (simulator.CLIFF_LAP, "CLIFF_LAP", "laps"),
                               (simulator.PASSABILITY, "PASSABILITY", "log-odds")):
        rows += [(f"{label}[{k}]", v, unit, "outcome / tyre model", "estimate") for k, v in table.items()]
    rows += [(f"GRIP[{cond}][{tyre}]", v, "-", "tyre grip relative to new soft, dry", "estimate")
             for cond, row in simulator.GRIP.items() for tyre, v in row.items()]
    s.table(["Parameter", "Value", "Units", "Used for", "Source"], rows)


def _write_experiment_sheets(wb):
    """Write one sheet per experiment from its saved ``reports/results`` JSON.

    Parameters
    ----------
    wb : openpyxl.Workbook
        Write-only workbook.
    """
    sol = _results("solver")
    s = _Sheet(wb, "Solver", [14, 22, 22, 22, 22])
    s.title("Euler and RK4 against analytical solutions")
    s.note("Position error after 12.8 s. Constant power (500 kW, no drag) and drag only (k = 0.6 kg/m).")
    titles = list(sol)
    s.header(["Time step (s)"] + [f"{t}: {m}" for t in titles for m in ("Euler error (m)", "RK4 error (m)")])
    for i, dt in enumerate(sol[titles[0]]["dt"]):
        s.row([dt] + [sol[t][m][i] for t in titles for m in ("euler", "rk4")])
    s.row(["Observed order"] + [round(sol[t][f"{m}_order"], 3) for t in titles for m in ("euler", "rk4")])

    s = _Sheet(wb, "Time step", [16, 18, 18, 16, 18, 18, 16])
    s.title("Braking-point deficit error against time step (2,000 scenarios, reference dt = 0.001 s)")
    for name, title in (("timestep", "Final model"), ("timestep_before_slipstream_fix",
                                                      "Before the slipstream fix (instant switch-off alongside)")):
        r = _results(name)
        s.section(title)
        s.header(["Time step (s)", "RK4 max error (s)", "RK4 mean error (s)", "RK4 run time (s)",
                  "Euler max error (s)", "Euler mean error (s)", "Euler run time (s)"])
        for i, dt in enumerate(r["dt"]):
            s.row([dt] + [r[m][k][i] for m in ("rk4", "euler") for k in ("max_error_s", "mean_error_s", "seconds")])
        s.row(["Observed order", round(r["rk4"]["observed_order"], 2), "", "", round(r["euler"]["observed_order"], 2)])

    quad = _results("quadrature")
    s = _Sheet(wb, "Quadrature", [20, 12, 22, 22])
    s.title("Quadrature rules and Monte Carlo (2,000 scenarios; reference: 200-node Gauss-Legendre)")
    s.header(["Rule", "n", "Evaluations per scenario", "Max absolute error in P"])
    for rule, r in quad["rules"].items():
        for n, ev, err in zip(r["n"], r["evaluations"], r["max_error"]):
            s.row([rule, n, ev, err])
    s.section("Monte Carlo replays (200 scenarios)")
    s.header(["Replays N", "RMS error", "Theory sqrt(p(1-p)/N)", ""])
    for row in zip(quad["monte_carlo"]["runs"], quad["monte_carlo"]["rms_error"], quad["monte_carlo"]["theory"]):
        s.row(list(row))
    s.section("Time per scenario")
    s.row(["Gauss-Legendre, 16 nodes (s)", quad["seconds_per_scenario"]["gauss_16"]])
    s.row(["Monte Carlo, 1,600 replays (s)", quad["seconds_per_scenario"]["monte_carlo_1600"]])

    sel = _results("selection")
    s = _Sheet(wb, "Network selection", [16, 12, 14, 9, 12, 14, 16, 14, 16, 16])
    s.title("Network size (single networks, validation set)")
    s.note(f"Rule: {sel['selection_rule']}. Lowest log loss: {sel['lowest_log_loss']}; threshold "
           f"{sel['threshold']:.4f}; selected: {sel['selected']}.")
    s.table(["Hidden layers", "Parameters", "FLOPs per prediction", "Epochs", "Fit time (s)",
             "Inference (us/prediction)", "Validation log loss", "Standard error", "Error vs truth (points)",
             "Default-scenario error (points)"],
            [(a["label"], a["parameters"], a["flops_per_prediction"], a["epochs"], round(a["fit_seconds"], 1),
              round(a["inference_us_per_prediction"], 3), round(a["val_log_loss"], 5), round(a["val_log_loss_se"], 5),
              round(100 * a["val_mae_vs_true"], 3), round(100 * a["default_scenario_error"], 2))
             for a in sel["architectures"]])
    s.section("Training-set size (selected network)")
    s.table(["Training attempts", "Validation log loss", "Error vs truth (points)"],
            [(d["train_size"], round(d["val_log_loss"], 5), round(100 * d["val_mae_vs_true"], 3))
             for d in sel["data_sizes"]])

    ab = _results("ablation")
    s = _Sheet(wb, "Ablation", [28, 8, 9, 16, 18, 22, 24])
    s.title("Physics-informed inputs and ensembling (validation set, five seeds)")
    keys = ("val_log_loss", "val_mae_vs_true", "default_scenario_error", "battery_axis_max_error")
    s.header(["Variant", "Seed", "Epochs", "Log loss", "Error vs truth (points)", "Default-scenario error (points)",
              "Worst battery-axis error (points)"])
    for variant in ("raw inputs only", "with energy-cover inputs"):
        for r in ab[variant]["runs"]:
            s.row([variant, r["seed"], r["epochs"], round(r[keys[0]], 5)] + [round(100 * r[k], 3) for k in keys[1:]])
    s.section("Single network (mean ± sd over seeds) and ensemble of five")
    s.header(["Variant", "", "", "Log loss", "Error vs truth (points)", "Default-scenario error (points)",
              "Worst battery-axis error (points)"])
    for variant in ("raw inputs only", "with energy-cover inputs"):
        single, ens = ab[variant]["single_model"], ab[variant]["ensemble_of_5"]
        s.row([f"{variant}: single", "", "", f"{single[keys[0]]['mean']:.4f} ± {single[keys[0]]['sd']:.4f}"]
              + [f"{100 * single[k]['mean']:.2f} ± {100 * single[k]['sd']:.2f}" for k in keys[1:]])
        s.row([f"{variant}: ensemble of 5", "", "", round(ens[keys[0]], 5)] + [round(100 * ens[k], 3) for k in keys[1:]])

    c = _results("comparison")
    s = _Sheet(wb, "Model comparison", [24, 16, 16, 16, 16, 16, 16, 16, 16, 16, 16])
    s.title(f"Test set ({c['test_attempts']:,} attempts)")
    cols = ["mae_vs_true", "max_error_vs_true", "log_loss", "brier", "roc_auc", "accuracy", "calibration_error",
            "fit_seconds", "predict_us_per_scenario", "model_bytes"]
    s.table(["Model", "Error vs truth", "Max error vs truth", "Log loss", "Brier", "ROC AUC", "Accuracy",
             "Calibration error", "Fit time (s)", "Predict (us/scenario)", "Model bytes"],
            [[name] + [m.get(k, "") for k in cols] for name, m in c["models"].items()])
    s.note("Error columns are probabilities (0.02 = 2 percentage points).")
    s.section("Mean error vs truth by operating region (points)")
    regions = list(next(iter(c["regions"].values())))
    s.table(["Region"] + list(c["regions"]), [[r] + [round(100 * c["regions"][m][r], 3) for m in c["regions"]]
                                              for r in regions])
    s.section("Error quantiles (points) and tyre-mismatch share of the worst errors")
    qkeys = list(next(iter(c["error_quantiles"].values())))
    s.table(["Statistic"] + list(c["error_quantiles"]),
            [[k] + [round(100 * c["error_quantiles"][m][k], 3) for m in c["error_quantiles"]] for k in qkeys])
    s.section("Extrapolation: straights of 1300-2000 m (mean error vs truth, points)")
    s.table(["Model", "Error"], [(m, round(100 * v, 3)) for m, v in c["extrapolation_1300_2000m"].items()])

    sens = _results("sensitivity")
    s = _Sheet(wb, "Sensitivity", [24, 12, 12, 16, 16, 16, 16])
    s.title("One input at a time around the default scenario")
    s.table(["Input", "Low setting", "High setting", "Network at low", "Network at high", "Truth at low",
             "Truth at high"],
            [(r["input"], r["low"], r["high"], round(r["network"][0], 4), round(r["network"][1], 4),
              round(r["exact"][0], 4), round(r["exact"][1], 4)) for r in sens["rows"]])
    s.section("Wet track, intermediates on both cars, attacker battery varied")
    s.table(["Attacker battery (%)", "Network", "Truth"],
            [(r["attacker_battery"], round(r["network"], 4), round(r["exact"], 4))
             for r in sens["wet_intermediates_battery_axis"]])
    w = sens["test_wet_straights_treaded_tyres"]
    s.note(f"Test set, wet straights on treaded tyres: {w['scenarios']:,} scenarios, mean signed error "
           f"{100 * w['mean_signed_error']:+.2f} points, mean absolute error {100 * w['mean_abs_error']:.2f} points.")

    perf = _results("performance")
    s = _Sheet(wb, "Performance", [34, 14, 16, 12])
    s.title("Profiling and performance")
    s.section("cProfile, 20,000 attempts: top functions by own time")
    s.table(["Function", "Calls", "Own time (s)", "Share"],
            [(p["function"], p["calls"], round(p["own_seconds"], 4), round(p["share"], 4))
             for p in perf["profile_top_functions"]])
    v = perf["vectorisation_200_scenarios"]
    s.section("Vectorisation (200 scenarios)")
    s.table(["One at a time (s)", "Together (s)", "Speed-up", ""],
            [(round(v["loop_seconds"], 3), round(v["vectorised_seconds"], 3), round(v["speedup"], 1), "")])
    s.section("Chunk size (40,000 attempts)")
    s.table(["Chunk", "Time (s)"], [(d["chunk"], round(d["seconds"], 2)) for d in perf["chunk_size_40k_scenarios"]])
    s.section("FLOPs and memory")
    s.table(["Quantity", "Value"],
            [(k, json.dumps(val) if isinstance(val, (dict, list)) else val) for k, val in perf["flops"].items()]
            + [(k, val) for k, val in perf["memory"].items()])


def _write_speed_trace(wb):
    """Write both cars' speed, position and energy every step along an 800 m straight (Figure A2's data).

    Parameters
    ----------
    wb : openpyxl.Workbook
        Write-only workbook.
    """
    s = _Sheet(wb, "Speed trace", [30, 10, 16, 16, 16, 16, 18, 18])
    s.title("Both cars along an 800 m straight (default scenario, and defender at 30 % battery)")
    s.header(["Scenario", "Time (s)", "Attacker position (m)", "Defender position (m)", "Attacker speed (km/h)",
              "Defender speed (km/h)", "Attacker energy used (MJ)", "Defender energy used (MJ)"])
    for label, scenario in (("Default (defender 60 %)", DEFAULTS),
                            ("Defender at 30 % battery", {**DEFAULTS, "defender_battery": 30.0})):
        trace = simulator.speed_trace(scenario).round(4)
        for r in trace.itertuples(index=False):
            s.row([label, r.t, r.attacker_x, r.defender_x, r.attacker_kph, r.defender_kph, r.attacker_mj,
                   r.defender_mj])


def write_workbook(path=OUT, n=300_000, log=print):
    """Build the attempts table and write the whole workbook to ``path``.

    Parameters
    ----------
    path : str or pathlib.Path, optional
    n : int, optional
        Number of simulated attempts to include.
    log : callable, optional

    Returns
    -------
    pathlib.Path
    """
    start = time.perf_counter()
    attempts = attempts_table(n, log=log)
    log(f"writing workbook ({len(attempts):,} rows)")
    wb = Workbook(write_only=True)
    # Arial throughout: replace the default font that unstyled cells use.
    wb._fonts = IndexedList([Font(name=FONT, sz=10, family=2)])
    _write_reference_sheets(wb, attempts)

    s = _Sheet(wb, "Attempts", [10, 11] + [14] * (len(attempts.columns) - 2))
    s.ws.freeze_panes = "C2"
    s.header(list(attempts.columns))
    for row in attempts.astype(object).itertuples(index=False):
        s.append([_native(v) for v in row])

    _write_speed_trace(wb)
    _write_experiment_sheets(wb)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log(f"saved {path} ({path.stat().st_size / 1e6:.1f} MB) in {time.perf_counter() - start:.0f}s")
    return path


def main(argv=None):
    """Command-line entry point: ``python -m overtake.export [--out PATH] [--samples N]``.

    Parameters
    ----------
    argv : list of str, optional
        Arguments without the program name; ``sys.argv[1:]`` if omitted.
    """
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--out", default=str(OUT), help="output .xlsx path")
    parser.add_argument("--samples", type=int, default=300_000, help="number of simulated attempts")
    args = parser.parse_args(argv)
    write_workbook(args.out, args.samples)


if __name__ == "__main__":
    main()
